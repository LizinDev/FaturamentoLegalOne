"""CLI, montagem da fila e exportacoes finais."""
import argparse
import csv
import datetime

import pytest
from openpyxl import Workbook, load_workbook

import config
import ledger as ledger_mod
import main
import planilha
from test_rodada import AutomadorFalso, achou, processo

# --- validacao de argumentos -------------------------------------------------

@pytest.mark.parametrize("texto, esperado", [("1", 1), ("500", 500)])
def test_inteiro_positivo_aceita(texto, esperado):
    assert main.inteiro_positivo(texto) == esperado


@pytest.mark.parametrize("texto", ["0", "-1", "abc", "1.5", ""])
def test_inteiro_positivo_recusa(texto):
    # `--max-cadastros 0` seria falso em Python e valeria como "sem cota": a
    # rodada iria ate o fim da planilha em vez de parar imediatamente.
    with pytest.raises(argparse.ArgumentTypeError):
        main.inteiro_positivo(texto)


def test_dia_iso_aceita_data_valida():
    assert main.dia_iso("2026-07-30") == "2026-07-30"


@pytest.mark.parametrize("texto", ["30/07/2026", "2026-13-01", "ontem", "2026-07-32"])
def test_dia_iso_recusa(texto):
    # O valor vira nome de arquivo e filtro LIKE do ledger.
    with pytest.raises(argparse.ArgumentTypeError):
        main.dia_iso(texto)


def test_limite_invalido_para_a_cli():
    with pytest.raises(SystemExit):
        main.argumentos(["--planilha", "x.xlsx", "--limite", "0"])


def test_padroes_da_cli():
    args = main.argumentos(["--planilha", "x.xlsx"])

    assert args.tarefa == config.PERFIL_PADRAO
    assert args.executar is False  # simulacao e o padrao
    assert args.limite is None


@pytest.mark.parametrize("segundos, esperado", [
    (0, "0m00s"), (59, "0m59s"), (75, "1m15s"), (3600, "1h00m"), (5432, "1h30m"),
])
def test_formatar_tempo(segundos, esperado):
    assert main._formatar_tempo(segundos) == esperado


# --- travas antes de comecar -------------------------------------------------

def test_planilha_e_obrigatoria():
    args = main.argumentos([])

    with pytest.raises(main.ErroDeUso, match="planilha"):
        main._modo_rodada(args)


def test_so_buscar_com_executar_e_recusado():
    # Marcaria no ledger como feito o que nunca foi cadastrado.
    args = main.argumentos(["--planilha", "x.xlsx", "--so-buscar", "--executar"])

    with pytest.raises(main.ErroDeUso, match="so-buscar"):
        main._modo_rodada(args)


def test_data_invalida():
    args = main.argumentos(["--planilha", "x.xlsx", "--data", "2026-07-30"])

    with pytest.raises(main.ErroDeUso, match="Data invalida"):
        main._data_da_tarefa(args)


def test_data_padrao_e_hoje():
    args = main.argumentos(["--planilha", "x.xlsx"])

    assert main._data_da_tarefa(args) == datetime.date.today().strftime("%d/%m/%Y")


def test_trava_de_planilha_por_perfil():
    perfil = config.PerfilTarefa("defesa-teste", "DEFESA FATURADA",
                                 dica_arquivo="defesa")
    args = main.argumentos(["--planilha", "C:/x/Faturamento.xlsx"])

    with pytest.raises(main.ErroDeUso, match="defesa"):
        main._conferir_planilha(args, perfil)

    # O nome certo passa, e --forcar-planilha manda no resto.
    main._conferir_planilha(main.argumentos(["--planilha", "C:/x/Defesa.xlsx"]), perfil)
    main._conferir_planilha(
        main.argumentos(["--planilha", "C:/x/Faturamento.xlsx", "--forcar-planilha"]),
        perfil,
    )


def test_perfil_sem_dica_nao_trava():
    # Dica vazia desliga a trava. O perfil e montado aqui, e nao tirado de
    # config.PERFIS, porque os perfis de verdade tem dica: se um dia deixarem
    # de ter, e a trava que some, nao este teste que quebra.
    perfil = config.PerfilTarefa("sem-dica", "TAREFA QUALQUER")
    args = main.argumentos(["--planilha", "qualquer.xlsx"])

    main._conferir_planilha(args, perfil)


def test_perfis_de_verdade_travam_a_planilha_trocada():
    """O par planilha/--tarefa e o erro caro: cadastra em lote a tarefa errada."""
    defesa_na_mao_do_faturamento = main.argumentos(
        ["--planilha", "C:/x/Defesa.xlsx", "--tarefa", "faturamento-final"]
    )
    faturamento_na_mao_da_defesa = main.argumentos(
        ["--planilha", "C:/x/Faturamento.xlsx", "--tarefa", "defesa-faturada"]
    )

    with pytest.raises(main.ErroDeUso, match="Faturamento"):
        main._conferir_planilha(defesa_na_mao_do_faturamento,
                                config.PERFIS["faturamento-final"])
    with pytest.raises(main.ErroDeUso, match="Defesa"):
        main._conferir_planilha(faturamento_na_mao_da_defesa,
                                config.PERFIS["defesa-faturada"])

    # E cada uma com a sua passa.
    main._conferir_planilha(
        main.argumentos(["--planilha", "C:/x/Faturamento 2026.xlsx"]),
        config.PERFIS["faturamento-final"],
    )
    main._conferir_planilha(
        main.argumentos(["--planilha", "C:/x/Defesa 2026.xlsx",
                         "--tarefa", "defesa-faturada"]),
        config.PERFIS["defesa-faturada"],
    )


def test_tarefa_auto_nao_exige_nome_de_arquivo():
    # No modo auto a garantia vem da celula de cada linha, que e mais forte do
    # que o nome do arquivo: nao ha planilha "da tarefa errada" para parear.
    args = main.argumentos(["--planilha", "C:/x/Planilha de Faturamento.xlsx",
                            "--tarefa", "auto"])

    main._conferir_planilha(args, config.PERFIL_AUTO)


def test_planilha_inexistente_vira_erro_de_uso(dados_tmp):
    args = main.argumentos(["--planilha", str(dados_tmp / "nao_existe.xlsx")])

    with pytest.raises(main.ErroDeUso):
        main._modo_rodada(args)


# --- montagem da fila --------------------------------------------------------

def _processos(*cnjs):
    return [processo(c) for c in cnjs]


def test_rodada_real_pula_o_que_ja_esta_no_ledger(registro, perfil):
    registro.registrar("A", perfil.descricao, ledger_mod.OK)
    registro.registrar("B", perfil.descricao, ledger_mod.NAO_ENCONTRADO)
    args = main.argumentos(["--planilha", "x.xlsx", "--executar"])

    _, fila, pular = main._montar_fila(args, registro, perfil, _processos("A", "B", "C"))

    assert [p.cnj for p in fila] == ["C"]
    assert pular == {("A", perfil.descricao), ("B", perfil.descricao)}


def test_retentar_refaz_erro_e_nao_encontrado(registro, perfil):
    registro.registrar("A", perfil.descricao, ledger_mod.OK)
    registro.registrar("B", perfil.descricao, ledger_mod.NAO_ENCONTRADO)
    registro.registrar("C", perfil.descricao, ledger_mod.ERRO)
    registro.registrar("D", perfil.descricao, ledger_mod.JA_EXISTIA)
    args = main.argumentos(["--planilha", "x.xlsx", "--executar", "--retentar"])

    _, fila, _ = main._montar_fila(args, registro, perfil,
                                   _processos("A", "B", "C", "D"))

    assert [p.cnj for p in fila] == ["B", "C"]


def test_simulacao_nao_pula_nada(registro, perfil):
    registro.registrar("A", perfil.descricao, ledger_mod.OK)
    args = main.argumentos(["--planilha", "x.xlsx"])

    _, fila, pular = main._montar_fila(args, registro, perfil, _processos("A", "B"))

    assert [p.cnj for p in fila] == ["A", "B"]
    assert pular == set()


def test_ledger_e_por_tarefa(registro, perfil):
    # O mesmo numero pode estar nas duas planilhas e precisar das duas tarefas.
    registro.registrar("A", "DEFESA FATURADA", ledger_mod.OK)
    args = main.argumentos(["--planilha", "x.xlsx", "--executar"])

    _, fila, _ = main._montar_fila(args, registro, perfil, _processos("A"))

    assert [p.cnj for p in fila] == ["A"]


def test_fila_auto_nao_pula_a_segunda_tarefa_do_mesmo_processo(registro):
    # O numero ja recebeu a defesa; continua devendo o faturamento final.
    registro.registrar("A", "DEFESA FATURADA", ledger_mod.OK)
    args = main.argumentos(["--planilha", "x.xlsx", "--tarefa", "auto", "--executar"])
    processos = [processo("A", tarefa="DEFESA FATURADA"),
                 processo("A", tarefa="FATURAMENTO FINAL")]

    _, fila, _ = main._montar_fila(args, registro, config.PERFIL_AUTO, processos)

    assert [p.tarefa for p in fila] == ["FATURAMENTO FINAL"]


def test_limite_corta_a_fila(registro, perfil):
    args = main.argumentos(["--planilha", "x.xlsx", "--limite", "2"])

    _, fila, _ = main._montar_fila(args, registro, perfil, _processos("A", "B", "C"))

    assert [p.cnj for p in fila] == ["A", "B"]


def test_processo_avulso_ignora_o_ledger_e_normaliza(registro, perfil, caplog):
    registro.registrar("0064904-50.2019.8.05.0001", perfil.descricao, ledger_mod.OK)
    args = main.argumentos([
        "--planilha", "x.xlsx", "--executar",
        # Como o usuario digita: com ponto no lugar do primeiro hifen.
        "--processo", "0064904.50.2019.8.05.0001", "0000000-00.0000.0.00.0000",
    ])

    processos, fila, _ = main._montar_fila(
        args, registro, perfil,
        _processos("0064904-50.2019.8.05.0001", "0000001-11.2025.8.05.0001"),
    )

    assert [p.cnj for p in fila] == ["0064904-50.2019.8.05.0001"]
    assert len(processos) == 1
    assert "0000000-00.0000.0.00.0000" in caplog.text  # avisa o que nao achou


# --- exportacoes -------------------------------------------------------------

def test_exportacao_que_falha_nao_leva_as_outras_junto(registro, dados_tmp, caplog):
    # No Windows, um relatorio.csv aberto no Excel devolve PermissionError.
    # Isto roda no finally da rodada, depois de horas de execucao.
    registro.registrar("A", "FATURAMENTO FINAL", ledger_mod.NAO_ENCONTRADO)

    def explode(*_):
        raise PermissionError("[Errno 13] arquivo aberto no Excel")

    registro.exportar_csv = explode

    main._exportar_finais(registro)

    assert "Nao consegui gravar o relatorio" in caplog.text
    # A lista de conferencia saiu mesmo assim.
    assert (dados_tmp / "nao_encontrados.csv").exists()


def test_exportar_finais_gera_a_planilha_do_dia(registro, dados_tmp):
    registro.con.execute(
        "INSERT INTO processos (cnj, tarefa, situacao, id_legalone, origem, quando) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        ("A", "FATURAMENTO FINAL", ledger_mod.OK, "111", "2026!L2",
         "2026-07-30T10:00:00"),
    )
    registro.con.commit()

    main._exportar_finais(registro, dias=["2026-07-30"])

    caminho = dados_tmp / "cadastrados_2026-07-30.xlsx"
    assert caminho.exists()
    ws = load_workbook(caminho).active
    assert ws["A2"].value == "A"
    assert ws["B2"].value == 111  # id numerico entra como numero


# --- modo relatorio ponta a ponta --------------------------------------------

def test_modo_relatorio_exporta_tudo(dados_tmp):
    with ledger_mod.Ledger(config.LEDGER_FILE) as led:
        led.con.execute(
            "INSERT INTO processos (cnj, tarefa, situacao, quando) VALUES (?,?,?,?)",
            ("A", "FATURAMENTO FINAL", ledger_mod.OK, "2026-07-30T10:00:00"),
        )
        led.con.commit()

    assert main.main(["--relatorio"]) == main.SAIDA_OK

    assert (dados_tmp / "relatorio.csv").exists()
    assert (dados_tmp / "nao_encontrados.csv").exists()
    assert (dados_tmp / "cadastrados_2026-07-30.xlsx").exists()


def test_modo_relatorio_de_um_dia_so(dados_tmp):
    with ledger_mod.Ledger(config.LEDGER_FILE) as led:
        led.con.executemany(
            "INSERT INTO processos (cnj, tarefa, situacao, quando) VALUES (?,?,?,?)",
            [("A", "FATURAMENTO FINAL", ledger_mod.OK, "2026-07-29T10:00:00"),
             ("B", "FATURAMENTO FINAL", ledger_mod.OK, "2026-07-30T10:00:00")],
        )
        led.con.commit()

    assert main.main(["--relatorio", "--dia", "2026-07-30"]) == main.SAIDA_OK

    assert (dados_tmp / "cadastrados_2026-07-30.xlsx").exists()
    assert not (dados_tmp / "cadastrados_2026-07-29.xlsx").exists()


def test_ledger_vazio_nao_quebra_o_relatorio(dados_tmp, caplog):
    assert main.main(["--relatorio"]) == main.SAIDA_OK

    assert "Nenhum cadastro registrado" in caplog.text


def test_sem_planilha_devolve_codigo_de_uso(dados_tmp):
    assert main.main([]) == main.SAIDA_USO


def test_fila_vazia_termina_sem_abrir_o_chrome(dados_tmp, monkeypatch, tmp_path):
    # Nenhum processo na planilha -> nem chega a conectar no Chrome.
    monkeypatch.setattr(planilha, "ler", lambda *a, **kw: [])
    def nao_deveria_conectar():
        raise AssertionError("nao era para tentar abrir o Chrome")
    monkeypatch.setattr(main.legalone, "conectar", nao_deveria_conectar)

    assert main.main(["--planilha", str(tmp_path / "Faturamento.xlsx")]) == main.SAIDA_OK


def test_chrome_fora_do_ar_devolve_codigo_de_rodada_abortada(
    dados_tmp, monkeypatch, tmp_path
):
    monkeypatch.setattr(planilha, "ler", lambda *a, **kw: [processo("A")])
    def sem_chrome():
        raise ConnectionRefusedError("porta 9222 fechada")
    monkeypatch.setattr(main.legalone, "conectar", sem_chrome)

    assert main.main(
        ["--planilha", str(tmp_path / "Faturamento.xlsx")]
    ) == main.SAIDA_ABORTADA


# --- rodada inteira ----------------------------------------------------------

ACHADO = "0000001-11.2025.8.05.0001"
INEXISTENTE = "0000002-22.2025.8.05.0001"
COM_PONTO = "0064904.50.2019.8.05.0001"          # o erro de digitacao da planilha
CORRIGIDO = "0064904-50.2019.8.05.0001"


def _linhas_csv(caminho) -> list[dict]:
    """Le um dos CSVs exportados (`;` e UTF-8 com BOM, para abrir no Excel)."""
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def _planilha_real(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    ws.append(["PROCESSO", "TIPO DE COBRANÇA", "STATUS LEGAL ONE"])
    ws.append([ACHADO, "ENCERRAMENTO FINAL", "Ativo"])
    ws.append([INEXISTENTE, "ACORDO", "Ativo"])
    ws.append([COM_PONTO, "ACORDO", "Baixado"])
    # O nome precisa combinar com a dica_arquivo do perfil padrao, senao a
    # trava do par planilha/--tarefa barra a rodada antes de ela comecar.
    caminho = tmp_path / "Faturamento.xlsx"
    wb.save(caminho)
    return caminho


def test_rodada_completa_da_planilha_aos_relatorios(dados_tmp, monkeypatch, tmp_path):
    """Planilha -> busca -> ledger -> os tres arquivos que o supervisor recebe."""
    arq = _planilha_real(tmp_path)
    automador = AutomadorFalso({ACHADO: achou("111"), CORRIGIDO: achou("222")})
    monkeypatch.setattr(main.legalone, "conectar", lambda: object())
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: automador)

    assert main.main(["--planilha", str(arq), "--executar"]) == main.SAIDA_OK

    # O numero com ponto foi corrigido antes de buscar.
    assert automador.buscados == [ACHADO, INEXISTENTE, CORRIGIDO]
    assert [c[0] for c in automador.cadastrados] == ["111", "222"]

    relatorio = {linha["PROCESSO"]: linha
                 for linha in _linhas_csv(dados_tmp / "relatorio.csv")}
    assert relatorio[ACHADO]["SITUACAO"] == "ok"
    assert relatorio[ACHADO]["TIPO_COBRANCA"] == "ENCERRAMENTO FINAL"
    assert relatorio[INEXISTENTE]["SITUACAO"] == "nao_encontrado"

    conferir = _linhas_csv(dados_tmp / "nao_encontrados.csv")
    assert [linha["PROCESSO"] for linha in conferir] == [INEXISTENTE]
    assert conferir[0]["ORIGEM"] == "2026!L3"
    # Rodada real nao encosta no arquivo da simulacao.
    assert not (dados_tmp / "nao_encontrados_simulacao.csv").exists()

    hoje = datetime.date.today().isoformat()
    ws = load_workbook(dados_tmp / f"cadastrados_{hoje}.xlsx").active
    assert [ws.cell(row=i, column=1).value for i in (2, 3)] == [ACHADO, CORRIGIDO]
    # Os valores fixos da tarefa sao repetidos em toda linha de proposito.
    assert ws["C2"].value == "FATURAMENTO FINAL"
    assert ws["D2"].value == "Cumprido"
    assert ws["F2"].value == "Heloiza Helena de Araujo"


def test_retomada_pula_o_que_ja_foi_e_preserva_a_lista_de_conferencia(
    dados_tmp, monkeypatch, tmp_path
):
    """A segunda rodada nao pode apagar o levantamento da primeira."""
    arq = _planilha_real(tmp_path)
    monkeypatch.setattr(main.legalone, "conectar", lambda: object())

    primeira = AutomadorFalso({ACHADO: achou("111"), CORRIGIDO: achou("222")})
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: primeira)
    main.main(["--planilha", str(arq), "--executar"])

    # Nada novo para fazer: os tres ja estao no ledger.
    segunda = AutomadorFalso({})
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: segunda)
    assert main.main(["--planilha", str(arq), "--executar"]) == main.SAIDA_OK

    assert segunda.buscados == []
    conferir = _linhas_csv(dados_tmp / "nao_encontrados.csv")
    assert [linha["PROCESSO"] for linha in conferir] == [INEXISTENTE]


def test_simulacao_nao_deixa_rastro_nenhum(dados_tmp, monkeypatch, tmp_path):
    arq = _planilha_real(tmp_path)
    automador = AutomadorFalso({ACHADO: achou("111"), CORRIGIDO: achou("222")})
    monkeypatch.setattr(main.legalone, "conectar", lambda: object())
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: automador)

    assert main.main(["--planilha", str(arq)]) == main.SAIDA_OK

    with ledger_mod.Ledger(config.LEDGER_FILE) as led:
        assert dict(led.resumo()) == {}
    # A lista da simulacao sai num arquivo proprio: uma simulacao de 20
    # processos nao pode apagar a lista de milhares das rodadas de verdade.
    conferir = _linhas_csv(dados_tmp / "nao_encontrados_simulacao.csv")
    assert [linha["PROCESSO"] for linha in conferir] == [INEXISTENTE]
    assert not (dados_tmp / "nao_encontrados.csv").exists()
    hoje = datetime.date.today().isoformat()
    assert not (dados_tmp / f"cadastrados_{hoje}.xlsx").exists()


def _planilha_mista(tmp_path):
    """Como a aba 2019-2020-2021: as duas tarefas na mesma aba, misturadas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "2019-2020-2021"
    ws.append(["PROCESSO", "TIPO DE COBRANÇA"])
    ws.append([ACHADO, "FATURAMENTO FINAL"])
    ws.append([CORRIGIDO, "DEFESA FATURADA"])
    # O mesmo processo nas duas etapas, e uma cobranca que nao e tarefa nenhuma.
    ws.append([CORRIGIDO, "FATURAMENTO FINAL"])
    ws.append([INEXISTENTE, "CONTESTAÇÃO"])
    caminho = tmp_path / "Planilha de Faturamento.xlsx"
    wb.save(caminho)
    return caminho


def test_rodada_auto_cadastra_a_tarefa_que_cada_linha_pede(
    dados_tmp, monkeypatch, tmp_path
):
    arq = _planilha_mista(tmp_path)
    automador = AutomadorFalso({ACHADO: achou("111"), CORRIGIDO: achou("222")})
    monkeypatch.setattr(main.legalone, "conectar", lambda: object())
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: automador)

    assert main.main(
        ["--planilha", str(arq), "--tarefa", "auto", "--executar"]
    ) == main.SAIDA_OK

    # A linha de CONTESTACAO nem entrou na fila: nao ha tarefa para ela.
    assert automador.buscados == [ACHADO, CORRIGIDO, CORRIGIDO]
    assert automador.tarefas == [
        "FATURAMENTO FINAL", "DEFESA FATURADA", "FATURAMENTO FINAL",
    ]

    tarefas_por_processo = {
        (linha["PROCESSO"], linha["TAREFA"]): linha["SITUACAO"]
        for linha in _linhas_csv(dados_tmp / "relatorio.csv")
    }
    assert tarefas_por_processo == {
        (ACHADO, "FATURAMENTO FINAL"): "ok",
        (CORRIGIDO, "DEFESA FATURADA"): "ok",
        (CORRIGIDO, "FATURAMENTO FINAL"): "ok",
    }

    hoje = datetime.date.today().isoformat()
    ws = load_workbook(dados_tmp / f"cadastrados_{hoje}.xlsx").active
    linhas = {(ws.cell(row=i, column=1).value, ws.cell(row=i, column=3).value)
              for i in range(2, 5)}
    assert linhas == {
        (ACHADO, "FATURAMENTO FINAL"),
        (CORRIGIDO, "DEFESA FATURADA"),
        (CORRIGIDO, "FATURAMENTO FINAL"),
    }


def test_simulacao_nao_apaga_a_lista_das_rodadas_de_verdade(
    dados_tmp, monkeypatch, tmp_path
):
    arq = _planilha_real(tmp_path)
    monkeypatch.setattr(main.legalone, "conectar", lambda: object())

    real = AutomadorFalso({ACHADO: achou("111"), CORRIGIDO: achou("222")})
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: real)
    main.main(["--planilha", str(arq), "--executar"])
    antes = (dados_tmp / "nao_encontrados.csv").read_bytes()

    # Uma simulacao curta depois, so com um processo que existe.
    simulado = AutomadorFalso({ACHADO: achou("111")})
    monkeypatch.setattr(main.legalone, "AutomadorLegalOne", lambda *a: simulado)
    main.main(["--planilha", str(arq), "--limite", "1"])

    assert (dados_tmp / "nao_encontrados.csv").read_bytes() == antes
