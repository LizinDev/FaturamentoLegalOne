"""Planilha Excel com os processos cadastrados num dia.

Separado do ledger.py de proposito: la e persistencia, aqui e apresentacao —
este arquivo e o que vai para o supervisor.
"""
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger(__name__)

CABECALHO = [
    ("PROCESSO", 30),
    ("ID LEGAL ONE", 14),
    ("TAREFA", 24),
    ("STATUS", 12),
    ("TIPO", 12),
    ("RESPONSÁVEL", 26),
    ("TIPO DE COBRANÇA", 34),
    ("STATUS NA PLANILHA", 20),
    ("ORIGEM NA PLANILHA", 30),
    ("CADASTRADO EM", 20),
]

_PREENCHIMENTO = PatternFill("solid", fgColor="1F3864")
_FONTE_CABECALHO = Font(bold=True, color="FFFFFF")


def _formatar_horario(iso: str) -> str:
    """2026-07-30T12:30:26 -> 30/07/2026 12:30:26"""
    try:
        data, hora = iso.split("T")
        ano, mes, dia = data.split("-")
        return f"{dia}/{mes}/{ano} {hora}"
    except (ValueError, AttributeError):
        return iso or ""


def gerar_planilha_do_dia(caminho: str | Path, dia: str, linhas: list[tuple]) -> int:
    """Grava o .xlsx do dia. `linhas` vem de Ledger.cadastrados_em()."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Cadastrados {dia}"

    for coluna, (titulo, largura) in enumerate(CABECALHO, 1):
        celula = ws.cell(row=1, column=coluna, value=titulo)
        celula.font = _FONTE_CABECALHO
        celula.fill = _PREENCHIMENTO
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(coluna)].width = largura

    for i, linha in enumerate(linhas, 2):
        cnj, tarefa, id_lo, tipo_cob, status_pl, origem, quando = linha
        # O ledger guarda so a descricao; o perfil devolve status, tipo e
        # responsavel. Um mesmo dia pode ter as duas tarefas.
        perfil = config.PERFIS_POR_DESCRICAO.get(tarefa)
        ws.cell(row=i, column=1, value=cnj)
        ws.cell(row=i, column=2, value=int(id_lo) if str(id_lo).isdigit() else id_lo)
        # Os valores da tarefa sao fixos por perfil; repeti-los em cada linha
        # deixa a planilha autoexplicativa para quem recebe e nao acompanhou.
        ws.cell(row=i, column=3, value=tarefa)
        ws.cell(row=i, column=4, value=perfil.status if perfil else "")
        ws.cell(row=i, column=5, value=perfil.tipo if perfil else "")
        ws.cell(row=i, column=6, value=perfil.responsavel_esperado if perfil else "")
        ws.cell(row=i, column=7, value=tipo_cob or "")
        ws.cell(row=i, column=8, value=status_pl or "")
        ws.cell(row=i, column=9, value=origem or "")
        ws.cell(row=i, column=10, value=_formatar_horario(quando))

    ws.freeze_panes = "A2"
    if linhas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{len(linhas) + 1}"

    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    logger.info("Planilha do dia salva: %s (%d processo[s])", caminho, len(linhas))
    return len(linhas)
