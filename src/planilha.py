"""Leitura da planilha de cobrancas -> lista de processos unicos."""
import collections
import dataclasses
import logging
import re
from collections.abc import Callable
from pathlib import Path

from openpyxl import load_workbook

import config

logger = logging.getLogger(__name__)

# NNNNNNN-DD.AAAA.J.TR.OOOO
RE_CNJ = re.compile(r"^\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}$")
# Mesma coisa, mas com ponto no lugar do hifen — erro de digitacao comum na
# planilha (ex.: 0064904.50.2019.8.05.0001). E recuperavel sem ambiguidade.
RE_CNJ_PONTO = re.compile(r"^(\d{7})\.(\d{2}\.\d{4}\.\d\.\d{2}\.\d{4})$")


@dataclasses.dataclass
class Processo:
    """Um processo unico da planilha, com a origem agregada."""

    cnj: str                 # numero normalizado, usado na busca
    cnj_original: str        # como estava na planilha
    formato_ok: bool
    # Descricao da tarefa a cadastrar neste processo. Vazia quando a tarefa vale
    # para a rodada toda; preenchida quando ela sai da propria linha.
    tarefa: str = ""
    tipos_cobranca: list[str] = dataclasses.field(default_factory=list)
    status_planilha: list[str] = dataclasses.field(default_factory=list)
    # Onde o processo aparece, no formato "aba!Lnn" — um processo repetido em
    # varias linhas/abas vira uma entrada so, com todas as origens.
    linhas: list[str] = dataclasses.field(default_factory=list)

    @property
    def origem(self) -> str:
        """Resumo de onde o processo aparece, para o relatorio."""
        return "; ".join(self.linhas)


def normalizar(numero: str) -> tuple[str, bool]:
    """Devolve (numero_normalizado, formato_ok).

    So corrige o caso do ponto no lugar do hifen. Numeros realmente quebrados
    passam adiante como estao: a busca vai falhar e eles aparecem no relatorio
    para conferencia manual, em vez de sumirem silenciosamente.
    """
    limpo = " ".join(str(numero).split())
    if RE_CNJ.match(limpo):
        return limpo, True
    m = RE_CNJ_PONTO.match(limpo)
    if m:
        return f"{m.group(1)}-{m.group(2)}", True
    return limpo, False


def _indices_cabecalho(linha) -> dict[str, int]:
    """Mapeia nome da coluna -> indice, a partir da linha de cabecalho."""
    indices = {}
    for i, valor in enumerate(linha):
        if valor is None:
            continue
        indices[" ".join(str(valor).split()).upper()] = i
    return indices


def _celula(linha, indices: dict[str, int], coluna: str) -> str:
    """Valor de uma coluna nomeada, ja normalizado ("" se ausente ou vazia)."""
    i = indices.get(coluna)
    if i is None or i >= len(linha) or linha[i] is None:
        return ""
    return " ".join(str(linha[i]).split())


def ler(
    caminho: str | Path,
    abas: list[str] | None = None,
    tipo_contem: str | None = None,
    status_planilha: str | None = None,
    tarefa_da_linha: Callable[[str], str | None] | None = None,
) -> list[Processo]:
    """Le a planilha e devolve os processos unicos, na ordem de aparicao.

    abas             — nomes de abas a considerar (None = todas)
    tipo_contem      — filtra por substring em TIPO DE COBRANCA (case-insensitive)
    status_planilha  — filtra por STATUS LEGAL ONE exato (ex.: "Ativo")
    tarefa_da_linha  — recebe o TIPO DE COBRANCA e devolve a descricao da tarefa
                       daquela linha, ou None para "nao reconheco isto" (a linha
                       e pulada). Com esta funcao a deduplicacao passa a ser por
                       (processo, tarefa): o mesmo numero que aparece como defesa
                       e como faturamento precisa das duas tarefas, e deduplicar
                       so pelo numero perderia uma delas.
    """
    caminho = Path(caminho)
    if not caminho.is_file():
        raise FileNotFoundError(f"Planilha nao encontrada: {caminho}")

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        alvo = abas or wb.sheetnames
        desconhecidas = [a for a in alvo if a not in wb.sheetnames]
        if desconhecidas:
            raise ValueError(
                f"Aba(s) inexistente(s): {desconhecidas}. "
                f"Disponiveis: {wb.sheetnames}"
            )

        # A chave e o par (numero, tarefa). Sem tarefa_da_linha a tarefa e "" em
        # todas as linhas e o par se comporta como a chave so pelo numero.
        por_chave: dict[tuple[str, str], Processo] = {}
        total_linhas = 0
        nao_reconhecidos: collections.Counter = collections.Counter()

        for nome in alvo:
            ws = wb[nome]
            indices: dict[str, int] = {}

            for n_linha, linha in enumerate(ws.iter_rows(values_only=True), 1):
                if n_linha == 1:
                    indices = _indices_cabecalho(linha)
                    if config.COLUNA_PROCESSO not in indices:
                        logger.warning(
                            "Aba %r sem coluna %r — ignorada",
                            nome, config.COLUNA_PROCESSO,
                        )
                        break
                    continue

                bruto = _celula(linha, indices, config.COLUNA_PROCESSO)
                if not bruto:
                    continue

                tipo = _celula(linha, indices, config.COLUNA_TIPO_COBRANCA)
                status = _celula(linha, indices, config.COLUNA_STATUS_LEGALONE)

                if tipo_contem and tipo_contem.upper() not in tipo.upper():
                    continue
                if status_planilha and status != status_planilha:
                    continue

                tarefa = ""
                if tarefa_da_linha is not None:
                    tarefa = tarefa_da_linha(tipo) or ""
                    if not tarefa:
                        # Pular e mais seguro do que chutar uma tarefa: a coluna
                        # e texto livre e as abas mais novas tem dezenas de
                        # variantes. Os valores saem no aviso do fim.
                        nao_reconhecidos[tipo or "(vazio)"] += 1
                        continue

                total_linhas += 1
                cnj, ok = normalizar(bruto)

                proc = por_chave.get((cnj, tarefa))
                if proc is None:
                    proc = Processo(cnj=cnj, cnj_original=bruto, formato_ok=ok,
                                    tarefa=tarefa)
                    por_chave[(cnj, tarefa)] = proc

                if tipo and tipo not in proc.tipos_cobranca:
                    proc.tipos_cobranca.append(tipo)
                if status and status not in proc.status_planilha:
                    proc.status_planilha.append(status)
                proc.linhas.append(f"{nome}!L{n_linha}")
    finally:
        wb.close()

    processos = list(por_chave.values())
    invalidos = sum(1 for p in processos if not p.formato_ok)
    logger.info(
        "Planilha lida: %d linha(s) -> %d processo(s) unico(s) (%d fora do padrao CNJ)",
        total_linhas, len(processos), invalidos,
    )
    if nao_reconhecidos:
        logger.warning(
            "%d linha(s) puladas por TIPO DE COBRANCA sem tarefa correspondente: %s",
            sum(nao_reconhecidos.values()),
            "; ".join(f"{valor!r} ({n})"
                      for valor, n in nao_reconhecidos.most_common()),
        )
    return processos
